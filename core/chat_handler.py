
import os
import json
import logging
import aiohttp
import mimetypes
import io
import re
import zoneinfo
from datetime import datetime, timedelta, time as dt_time
from PIL import Image
from google import genai
from google.genai import types
from pydantic import BaseModel, Field
from typing import Optional, List

logger = logging.getLogger("ChatHandler")


class ReactionAction(BaseModel):
    self_emojis: List[str] = Field(
        default_factory=list, 
        description="List of unicode emoji characters to add directly to the bot's own message as reactions."
    )
    user_emojis: List[str] = Field(
        default_factory=list, 
        description="List of unicode emoji characters to add directly to the user's incoming message as reactions."
    )

class ChannelWatchAction(BaseModel):
    channel_id: int = Field(
        ..., 
        description="The snowflake ID of the Discord channel to target."
    )
    action: str = Field(
        ..., 
        description="Either 'watch' to begin active channel monitoring, or 'unwatch' to stop."
    )

class TypoCorrectionAction(BaseModel):
    original: str = Field(
        ..., 
        description="A minor typo word or spelling slip to intentionally send first."
    )
    corrected: str = Field(
        ..., 
        description="The corrected word that will dynamically overwrite the typo via edit after 2.5 seconds."
    )

class MemoryAction(BaseModel):
    operation: str = Field(
        ..., 
        description="Either 'learn' to save a new observation fact or 'forget' to remove a fact."
    )
    tier: str = Field(
        ..., 
        description="Either 'user' (user memories), 'server' (server lore), or 'global' (shared DB)."
    )
    fact: str = Field(
        ..., 
        description="The highly specific, concise fact to remember or forget."
    )

class PollAction(BaseModel):
    question: str = Field(
        ..., 
        description="The main question prompt for the native Discord poll."
    )
    options: List[str] = Field(
        ..., 
        description="A list of 2 to 10 choices or answers for the poll."
    )
    duration_hours: int = Field(
        default=24, 
        description="The lifespan of the poll in hours."
    )

class ButtonComponent(BaseModel):
    label: str = Field(
        ..., 
        description="Text displayed on the button (max 80 chars)."
    )
    style: str = Field(
        default="secondary", 
        description="Preset color style: primary, secondary, success, danger."
    )
    emoji: Optional[str] = Field(
        default=None, 
        description="Highly recommended. A relevant visual unicode emoji symbol."
    )

class SelectOption(BaseModel):
    label: str = Field(
        ..., 
        description="The visible text label of this option (max 100 chars)."
    )
    description: Optional[str] = Field(
        default=None, 
        description="Required. A detailed, context-specific description explaining what choosing this option will trigger (max 100 chars)."
    )
    emoji: Optional[str] = Field(
        default=None, 
        description="Required. A relevant visual unicode or custom emoji symbol matching the option theme."
    )

class DropdownComponent(BaseModel):
    type: str = Field(
        ..., 
        description="Type of dropdown: 'string' (custom choices), 'user', 'role', 'channel', 'mentionable'."
    )
    placeholder: str = Field(
        ..., 
        description="Helper text in empty dropdown (max 150 characters)."
    )
    options: List[SelectOption] = Field(
        default_factory=list, 
        description="Selection options. Only required and used when type is 'string'."
    )

class ModalField(BaseModel):
    label: str = Field(
        ..., 
        description="The prompt heading text above this field's input box."
    )
    style: str = Field(
        default="short", 
        description="Style of input: 'short', 'long', 'user_select', 'role_select', 'channel_select', 'mentionable_select', or 'select_string'."
    )
    description: Optional[str] = Field(
        default=None, 
        description="Optional placeholder text inside the input box."
    )
    options: Optional[List[SelectOption]] = Field(
        default=None, 
        description="Predefined list of choices. Only used and required when style is 'select_string'. Each option can have a label, description, and emoji."
    )

class ModalComponent(BaseModel):
    button_label: str = Field(
        ..., 
        description="The label of the button that will launch the modal popup."
    )
    modal_title: str = Field(
        ..., 
        description="The title of the modal popup window (max 45 chars)."
    )
    fields: List[ModalField] = Field(
        ..., 
        description="The list of form fields inside this popup modal."
    )

class InteractiveComponents(BaseModel):
    buttons: List[ButtonComponent] = Field(
        default_factory=list, 
        description="List of clickable option buttons (max 5 per row, max 25 total)."
    )
    dropdowns: List[DropdownComponent] = Field(
        default_factory=list, 
        description="List of interactive selection dropdown menus (max 1 per row)."
    )
    modals: List[ModalComponent] = Field(
        default_factory=list, 
        description="List of buttons that launch custom data collection forms."
    )

class CollectorAction(BaseModel):
    title: str = Field(
        ..., 
        description="Title of the survey or collector session."
    )
    duration_seconds: int = Field(
        default=60, 
        description="Active collection window in seconds."
    )
    is_anonymous: bool = Field(
        default=True, 
        description="True if submissions should be anonymized, False for public logging."
    )



class ConversationalResponse(BaseModel):
    content: str = Field(
        ..., 
        description="The main text response. MUST be written mostly in lowercase, using casual punctuation (no ending periods on single sentences), and zero robotic layouts. Do NOT write raw LaTeX math symbols (such as $...$ or $$...$$) inside this text; use simple plain text, code blocks, or unicode characters."
    )
    follow_ups: List[str] = Field(
        default_factory=list, 
        description="Optional list of consecutive messages to send immediately after the main content."
    )
    reactions: Optional[ReactionAction] = Field(
        default=None, 
        description="Reactions to apply to standard messages."
    )
    channel_watch: Optional[ChannelWatchAction] = Field(
        default=None, 
        description="Active channel watch or unwatch controls."
    )
    typo_correction: Optional[TypoCorrectionAction] = Field(
        default=None, 
        description="Typo and edit self-corrections."
    )
    reset_history: bool = Field(
        default=False, 
        description="True if the conversational tracker cache for this channel should be wiped."
    )
    memory_actions: List[MemoryAction] = Field(
        default_factory=list, 
        description="Observations to learn or forget in the long-term databases."
    )
    poll: Optional[PollAction] = Field(
        default=None, 
        description="Native Discord poll parameters."
    )
    components: Optional[InteractiveComponents] = Field(
        default=None, 
        description="Legacy interactive UI components (buttons, dropdowns, modal triggers)."
    )
    collector: Optional[CollectorAction] = Field(
        default=None, 
        description="Multi-user response tracking collector."
    )
    message_builder_v2_dsl: Optional[str] = Field(
        default=None, 
        description="Complete, valid Python DSL layout string representing components V2. Rules: 1. Always use Container(...) as the root component (Never use Layout). 2. Section() objects MUST have exactly 1 Button accessory passed as a keyword-only argument (e.g. accessory=Button(...)). 3. A Section can only hold between 1 and 3 TextDisplay objects on the left column (In components.0.components: Must be between 1 and 3 in length). If you only want text with no accessory, use standard Container(TextDisplay(...)) instead of a Section!"
    )
    rendered_latex_steps: Optional[List[str]] = Field(
        default=None, 
        description="A list of pure, clean LaTeX mathematical formulas to render sequentially on the visual math card (e.g., ['\\\\int_0^\\\\infty \\\\frac{x^3}{e^x-1} dx', '= \\\\frac{\\\\pi^4}{15}']). Each entry MUST contain ONLY math formulas, with absolutely zero natural language words, sentences, or conversational dialogue."
    )



class ChatHandler:
    def __init__(self, api_key: str, system_prompt_path: str, model_name: str = None):
        self.client = genai.Client(api_key=api_key)
        
        self.premium_model = os.getenv("GEMINI_MODEL_PREMIUM", "gemini-3-flash-preview")
        self.fallback_model = os.getenv("GEMINI_MODEL_TEXT", "gemma-4-31b-it")
        
        self.pt_zone = zoneinfo.ZoneInfo("America/Los_Angeles")
        self.premium_cooldown_until = None
        
        logger.info(f"ChatHandler initialized. Premium: {self.premium_model} | Fallback: {self.fallback_model}")
        
        try:
            with open(system_prompt_path, 'r', encoding='utf-8') as file:
                self.system_instruction = file.read()
        except FileNotFoundError:
            logger.error(f"Could not find system prompt at {system_prompt_path}")
            self.system_instruction = "You are a helpful friend on Discord."

        artist_prompt_path = os.path.join("config", "image_prompt_system.md")
        try:
            with open(artist_prompt_path, 'r', encoding='utf-8') as file:
                self.artist_system_instruction = file.read()
        except FileNotFoundError:
            logger.error(f"Could not find specialized image system prompt at {artist_prompt_path}")
            self.artist_system_instruction = "Translate the user request into a highly detailed Stable Diffusion prompt."

    def _check_and_reset_quota(self):
        now_pt = datetime.now(self.pt_zone)
        if self.premium_cooldown_until and now_pt >= self.premium_cooldown_until:
            logger.info("API daily reset threshold passed. Gemini 3 Premium re-enabled.")
            self.premium_cooldown_until = None
        return now_pt

    async def generate_artist_stream(self, raw_prompt: str, context_history: str = "", base_image_bytes: bytes = None):
        now_pt = self._check_and_reset_quota()
        active_model = self.premium_model if not self.premium_cooldown_until else self.fallback_model
        use_thinking = (active_model == self.premium_model)
        
        contents = []
        if base_image_bytes:
            try:
                img = Image.open(io.BytesIO(base_image_bytes))
                contents.append(img)
            except Exception as e:
                logger.error(f"Failed to parse base image for Artist: {e}")
                
        contents.append(f"--- CONVERSATIONAL REFERENCE HISTORY ---\n{context_history}\n\n")
        contents.append(f"--- TARGET USER REQUEST TO EXPAND/EDIT ---\n{raw_prompt}")

        if use_thinking:
            config = types.GenerateContentConfig(
                system_instruction=self.artist_system_instruction,
                temperature=0.7,
                thinking_config=types.ThinkingConfig(
                    thinking_level="HIGH",
                    include_thoughts=True
                )
            )
        else:
            config = types.GenerateContentConfig(
                system_instruction=self.artist_system_instruction,
                temperature=0.7
            )

        try:
            return await self.client.aio.models.generate_content_stream(
                model=active_model,
                contents=contents,
                config=config
            )
        except Exception as e:
            error_str = str(e).lower()
            if "429" in error_str or "exhausted" in error_str or "quota" in error_str:
                if self.premium_model == self.premium_model:
                    logger.warning("Premium Rate limit hit during Artist stream init! Locking and falling back.")
                    tomorrow_pt = now_pt + timedelta(days=1)
                    self.premium_cooldown_until = datetime.combine(tomorrow_pt.date(), dt_time(0, 0, 0), tzinfo=self.pt_zone)
                    
                    fallback_config = types.GenerateContentConfig(
                        system_instruction=self.artist_system_instruction,
                        temperature=0.7
                    )
                    return await self.client.aio.models.generate_content_stream(
                        model=self.fallback_model,
                        contents=contents,
                        config=fallback_config
                    )
            raise e

    async def _download_attachment(self, url: str) -> bytes:
        async with aiohttp.ClientSession() as session:
            async with session.get(url) as response:
                if response.status == 200:
                    return await response.read()
                return b""

    def _should_use_search(self, message_content: str, channel_history: str = "") -> bool:
        content = (message_content + " " + channel_history).lower()
        search_keywords = [
            "weather", "temperature", "news", "current", "latest", "time in", "date", 
            "who is", "what is the price of", "stock", "google", "search", "lookup"
        ]
        return any(kw in content for kw in search_keywords)

    def _select_thinking_level(self, message_content: str, channel_history: str) -> str:
        content = message_content.lower().strip()
        
        ui_triggers = [
            "interacted with a ui component",
            "active v2 component",
            "just reacted to your message",
            "removed their reaction",
            "submitted form",
            "clicked button",
            "selected option",
            "selected user",
            "selected role",
            "selected channel"
        ]
        
        if any(trigger in content for trigger in ui_triggers):
            high_triggers = [
                "proof", "paradox", "riddle", "algorithm", "mathematical", "complex", "debug", 
                "architecture", "scale", "optimize", "system design", "write code", "implement",
                "explain how", "solve", "math", "analysis", "deduce"
            ]
            if not any(t in content for t in high_triggers):
                return "NONE"

        if message_content.startswith("[System") or "system prompt" in content:
            high_triggers = [
                "proof", "paradox", "riddle", "algorithm", "mathematical", "complex", "debug", 
                "architecture", "scale", "optimize", "system design", "write code", "implement",
                "explain how", "solve", "math", "analysis", "deduce"
            ]
            if any(t in content for t in high_triggers):
                return "HIGH"
            return "NONE"

        clean_content = re.sub(r'<@\d+>', '', content).strip()
        clean_content = re.sub(r'[^\w\s]', '', clean_content).strip()
        
        casual_greetings_and_replies = {
            "hi", "hello", "hey", "yo", "sup", "greetings", "howdy", "morning", "afternoon", "evening",
            "thanks", "thank you", "ty", "thx", "ok", "okay", "cool", "nice", "awesome", "lol", "lmao",
            "haha", "hahaha", "yes", "no", "yep", "nope", "sure", "fine", "good", "bad", "sad", "bruh", "wsp"
        }
        
        words = clean_content.split()
        if not words or (len(words) <= 3 and all(w in casual_greetings_and_replies for w in words)):
            return "NONE"
            
        if len(clean_content) < 15:
            high_triggers = [
                "proof", "paradox", "riddle", "algorithm", "mathematical", "complex", "debug", 
                "architecture", "scale", "optimize", "system design", "write code", "implement",
                "explain how", "solve", "math", "analysis", "deduce"
            ]
            if not any(t in clean_content for t in high_triggers):
                return "NONE"

        high_triggers = [
            "proof", "paradox", "riddle", "algorithm", "mathematical", "complex", "debug", 
            "architecture", "scale", "optimize", "system design", "write code", "implement",
            "explain how", "solve", "math", "analysis", "deduce"
        ]
        if any(t in clean_content for t in high_triggers):
            return "HIGH"
            
        return "NONE"

    def _sanitize_mime_type_and_data(self, mime_type: str, filename: str, file_bytes: bytes) -> tuple[str, bytes] | tuple[None, None]:
        if not mime_type:
            return None, None
            
        mime_type = mime_type.lower().split(';')[0].strip()
        
        supported_images = ["image/png", "image/jpeg", "image/jpg", "image/webp", "image/heic", "image/heif"]
        if mime_type in supported_images or (mime_type.startswith("image/") and mime_type != "image/gif"):
            return mime_type if mime_type in supported_images else "image/png", file_bytes
        
        supported_audio = ["audio/wav", "audio/mp3", "audio/aiff", "audio/aac", "audio/ogg", "audio/flac", "audio/mpeg", "audio/x-m4a"]
        if mime_type in supported_audio or mime_type.startswith("audio/"):
            return mime_type if mime_type in supported_audio else "audio/wav", file_bytes
            
        supported_video = ["video/mp4", "video/mpeg", "video/mov", "video/avi", "video/flv", "video/mpg", "video/webm", "video/wmv", "video/3gpp", "video/quicktime"]
        if mime_type in supported_video or mime_type.startswith("video/"):
            return mime_type if mime_type in supported_video else "video/mp4", file_bytes
            
        if mime_type == "application/pdf" or filename.lower().endswith(".pdf"):
            return "application/pdf", file_bytes

        text_extensions = (
            '.py', '.json', '.md', '.txt', '.js', '.ts', '.sh', '.css', '.html', 
            '.xml', '.yaml', '.yml', '.c', '.cpp', '.h', '.java', '.go', '.rs', 
            '.sql', '.bat', '.ps1', '.ini', '.conf', '.env', '.log'
        )
        if mime_type.startswith("text/") or filename.lower().endswith(text_extensions):
            return "text/plain", file_bytes

        if filename.lower().endswith(".docx"):
            try:
                import docx
                doc = docx.Document(io.BytesIO(file_bytes))
                text = "\n".join([p.text for p in doc.paragraphs])
                return "text/plain", text.encode('utf-8')
            except Exception as e:
                logger.warning(f"Docx extraction failed for {filename}: {e}")

        if filename.lower().endswith(".xlsx"):
            try:
                import openpyxl
                wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
                lines = []
                for sheet in wb.worksheets:
                    lines.append(f"--- Sheet: {sheet.title} ---")
                    for row in sheet.iter_rows(values_only=True):
                        lines.append("\t".join([str(c) if c is not None else "" for c in row]))
                return "text/plain", "\n".join(lines).encode('utf-8')
            except Exception as e:
                logger.warning(f"Xlsx extraction failed for {filename}: {e}")
            
        return None, None

    async def generate_reply_stream(self, message_content: str, channel_history: str, attachments: list, user_display_name: str, user_memory: dict, server_context: str, scraped_pages: list[str] = None, user_status=None, thinking_level="NONE", is_dm=False, active_config=None, bot_instance=None):
        now_pt = self._check_and_reset_quota()
        
        if self.premium_cooldown_until:
            active_model = self.fallback_model
        else:
            active_model = self.premium_model

        if active_config is None: active_config = {}

        custom_prompt = active_config.get("system_prompt", "").strip()
        base_sys_prompt = custom_prompt if custom_prompt else self.system_instruction

        math_system_guardrails = (
            "\n\n=== MATHEMATICAL LATEX & INTEGRATION RULES (MANDATORY) ===\n"
            "1. DISCORD LATEX LIMITATION: Discord does NOT natively render raw LaTeX formatting "
            "in text chat (e.g., surrounds like $...$ or $$...$$ will display as plain, unformatted text). "
            "Therefore, you are STRICTLY FORBIDDEN from writing raw LaTeX formulas, equations, "
            "or mathematical delimiters inside your conversational 'content' string. Express math "
            "using standard plain text, basic unicode symbols (e.g. ∫, π, ², oo), or clean code blocks. "
            "Leave complex LaTeX typesetting strictly to the generated visual math cards.\n"
            "2. SYMPY INTEGRATION METHOD: SymPy cannot solve complex definite integrals involving "
            "Euler exponents and polylogarithms directly. When requested to solve a definite integral, "
            "do NOT pass the raw integrate(...) command directly. Instead, logically deconstruct the "
            "definite integral into its symbolic components (such as Γ(s), ζ(s), and their derivatives) "
            "using your thoughts, evaluate those components using SymPy tools, and then programmatically "
            "re-assemble the exact closed-form expression.\n"
            "3. COMPUTATION VS. SIGHT: Use the `execute_math_evaluation` tool ONLY when you need to perform "
            "raw algebraic solving, calculus derivatives/integrals, simplification, or coordinate plotting. "
            "For multiple-choice graph matching, visual pattern recognition, geometric shape identification, or reading coordinate grids from images, "
            "do NOT call the math tool. Solve these problems directly using your multimodal sight, work out the parameters analytically, "
            "and present your explanation in clean plain text.\n"
            "4. REQUIRED PLOT ACTION: If the user explicitly asks you to 'plot', 'graph', or 'draw' a mathematical "
            "equation, you are REQUIRED to trigger the `execute_math_evaluation` tool to generate the visual graph, "
            "even if you can solve the algebra in your head.\n"
            "5. LATEX PARSING CAPABILITY: Note that your `execute_math_evaluation` tool is fully equipped to parse "
            "raw LaTeX math strings. If you need to evaluate mathematical terms with standard LaTeX formatting "
            "(e.g. \\frac, \\int, \\infty, exponents, etc.) submitted by a user, you can pass that raw LaTeX string "
            "directly as the `query` argument."
        )
        base_sys_prompt += math_system_guardrails

        if is_dm:
            base_sys_prompt = re.sub(r'<!-- THREAD_INSTRUCTIONS_START -->.*?<!-- THREAD_INSTRUCTIONS_END -->', '', base_sys_prompt, flags=re.DOTALL)
            base_sys_prompt = base_sys_prompt.replace("deploy a side-thread using the `[THREAD]` tag, and then deliver your code modules **sequentially and modularly**.", "deliver your code modules sequentially in this DM.")
            base_sys_prompt += "\n\nCRITICAL DM RULE: You are currently chatting in Direct Messages (DMs). There are NO threads in DMs. You are STRICTLY FORBIDDEN from using `[THREAD]` tags, attempting to spawn threads, or referencing thread creation. Treat all exploratory requests inline in this DM."
            base_sys_prompt += "\n\nCRITICAL DM RULE 2: Since you are currently chatting in Direct Messages (DMs), there are no server-side roles, channels, or mentionables. You are STRICTLY FORBIDDEN from instantiating or using RoleSelect, ChannelSelect, or MentionableSelect components inside your visual layouts or modal popups."

        disc_tools = active_config.get("discord_tools", [])

        if "Server Emojis" not in disc_tools:
            base_sys_prompt += "\n\nCRITICAL EMOJI RESTRICTION: You are STRICTLY FORBIDDEN from generating or writing any custom server emojis (do not use the <:name:id> or <a:name:id> format). Write using standard text or unicode emojis only."
        if "Unicode Emojis" not in disc_tools:
            base_sys_prompt += "\n\nCRITICAL EMOJI RESTRICTION: You are STRICTLY FORBIDDEN from outputting standard unicode emojis (e.g. 🙂, 🔥, 😂, 👀, 💀) in your messages. Express yourself purely through text."

        parts = []
        status_section = f"--- CURRENT STATUS & ACTIVITY FOR {user_display_name} ---\n{user_status}\n\n" if user_status else ""
        scraped_section = "\n\n--- SCRAPED WEBPAGE CONTENT ---\n" + "\n\n---\n\n".join(scraped_pages) if scraped_pages else ""
        
        context_prompt = (
            f"--- SERVER ENVIRONMENT DATA ---\n"
            f"{server_context}\n\n"
            f"--- LONG TERM MEMORY FILE FOR {user_display_name} ---\n"
            f"{json.dumps(user_memory, indent=2)}\n\n"
            f"--- MULTI-USER CONVERSATION TRANSCRIPT ---\n"
            f"{channel_history}\n\n"
            f"{status_section}"
            f"{scraped_section}\n\n"
            f"--- NEW MESSAGE FROM {user_display_name} ---\n"
            f"{message_content}\n\n"
            f"--- YOUR TURN ---\n"
            f"Generate ONLY your single response. Do not continue the transcript. Talk directly to {user_display_name}."
        )
        parts.append(context_prompt)

        for attachment in attachments:
            raw_mime = attachment.content_type
            if not raw_mime:
                raw_mime, _ = mimetypes.guess_type(attachment.filename)
            if not raw_mime:
                continue
                
            file_bytes = await self._download_attachment(attachment.url)
            if not file_bytes:
                continue
                
            mime_type, sanitized_bytes = self._sanitize_mime_type_and_data(raw_mime, attachment.filename, file_bytes)
            if mime_type and sanitized_bytes:
                parts.append(types.Part.from_bytes(data=sanitized_bytes, mime_type=mime_type))

        thinking_config = None
        if thinking_level in ("HIGH", "MINIMAL"):
            thinking_config = types.ThinkingConfig(
                thinking_level=thinking_level,
                include_thoughts=True
            )

        tools_list = []
        if bot_instance:
            from core.tools import ToolSuite
            suite = ToolSuite(bot_instance)
            
            enabled_sys = active_config.get("system_tools", [])
            if "Memory Journals" in enabled_sys:
                tools_list.append(suite.save_memory_fact)
                tools_list.append(suite.forget_memory_fact)
            if "Google Search" in enabled_sys:
                tools_list.append(suite.web_search)
            if "URL Content" in enabled_sys:
                tools_list.append(suite.web_scrape)
            if "Generate Images" in enabled_sys:
                tools_list.append(suite.generate_image)
            
            disable_math_tool = False
            if attachments:
                visual_cues = ["graph", "plot", "midline", "amplitude", "period", "shown below", "coordinate", "which equation", "figure", "chart", "diagram", "triangle"]
                content_lower = message_content.lower()
                if any(cue in content_lower for cue in visual_cues):
                    disable_math_tool = True
                    logger.info("Dynamic Math Tool Omission: Omit math execution tool because prompt implies visual graph reading or geometry analysis.")

            if "Code Execution" in enabled_sys and not disable_math_tool:
                tools_list.append(suite.execute_math_evaluation)
                
            if "Threads" in disc_tools:
                tools_list.append(suite.create_thread)
            
            tools_list.append(suite.watch_channel)

        mode_mapping = {
            "Auto": "AUTO",
            "Forced": "ANY",
            "Off": "NONE"
        }
        selected_mode = mode_mapping.get(active_config.get("tool_mode", "Auto"), "AUTO")
        tool_config_obj = types.ToolConfig(
            function_calling_config=types.FunctionCallingConfig(mode=selected_mode)
        )
        
        config = types.GenerateContentConfig(
            system_instruction=base_sys_prompt,
            temperature=0.7,
            response_mime_type="application/json",
            response_schema=ConversationalResponse,
            thinking_config=thinking_config,
            tools=tools_list if tools_list else None,
            tool_config=tool_config_obj,
            automatic_function_calling=types.AutomaticFunctionCallingConfig(disable=True)
        )

        try:
            return await self.client.aio.models.generate_content_stream(
                model=active_model,
                contents=parts,
                config=config
            )
        except Exception as e:
            error_str = str(e).lower()
            if "429" in error_str or "exhausted" in error_str or "quota" in error_str:
                if active_model == self.premium_model:
                    logger.warning("Premium model quota exhausted during stream initialization! Locking premium and falling back.")
                    tomorrow_pt = now_pt + timedelta(days=1)
                    self.premium_cooldown_until = datetime.combine(tomorrow_pt.date(), dt_time(0, 0, 0), tzinfo=self.pt_zone)
                    
                    fallback_config = types.GenerateContentConfig(
                        system_instruction=base_sys_prompt, 
                        temperature=0.7,
                        response_mime_type="application/json",
                        response_schema=ConversationalResponse,
                        tools=tools_list if tools_list else None,
                        tool_config=tool_config_obj,
                        automatic_function_calling=types.AutomaticFunctionCallingConfig(disable=True)
                    )
                    return await self.client.aio.models.generate_content_stream(
                        model=self.fallback_model, contents=parts, config=fallback_config
                    )
            raise e